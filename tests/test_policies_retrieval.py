import io

import docx

from policies.parser import parse_policy
from retrieval.hybrid import search

MD = b"""# Widget Policy
Policy ID: WID
Domain: logging

## 1 Intro
Background only.

## 2.1 Logs
Widgets must emit audit logs. (risk: high; evidence: logging)
Widgets should be blue.
"""


def test_markdown_sections_requirements_and_tags():
    p = parse_policy(MD, "widget.md")
    assert (p.key, p.title, p.domain) == ("WID", "Widget Policy", "logging")
    assert [s.number for s in p.sections] == ["1", "2.1"]
    (req,) = p.sections[1].requirements
    assert req.code == "WID-2.1.1"
    assert req.risk == "high" and req.evidence_types == ["logging"]
    assert "(risk" not in req.text and "(risk" not in p.sections[1].text
    assert p.sections[0].requirements == []


def test_inferred_risk_and_types_without_tag():
    p = parse_policy(b"# P\n## 1 S\nCustomer data must be encrypted at rest.\n", "p.md")
    (req,) = p.sections[0].requirements
    assert req.risk == "high" and "encryption" in req.evidence_types


def test_html_and_docx_formats():
    html = b"<h1>Html Policy</h1><h2>1 Access</h2><p>Admins must use MFA.</p>"
    p = parse_policy(html, "h.html")
    assert p.title == "Html Policy" and p.sections[0].requirements[0].text == "Admins must use MFA."

    d = docx.Document()
    d.add_heading("Docx Policy", 1)
    d.add_heading("1 Retention", 2)
    d.add_paragraph("Records must be purged after 7 years.")
    buf = io.BytesIO()
    d.save(buf)
    p = parse_policy(buf.getvalue(), "d.docx")
    assert p.title == "Docx Policy"
    assert p.sections[0].requirements[0].evidence_types == ["retention"]


def test_hybrid_search_returns_cited_sections(session):
    hits = search(session, "We are deploying a new API that stores customer information", k=10)
    citations = {h.section.citation.split(" ")[0] for h in hits}
    assert {"DATA-CLASS", "DATA-RET", "SEC-ACCESS"} <= citations
    assert all("§" in h.section.citation for h in hits)


def test_keyword_query_ranks_exact_section_first(session):
    (top, *_) = search(session, "multi-factor authentication for administrative access", k=3)
    assert top.section.citation == "SEC-ACCESS §2.2"


def test_search_falls_back_to_keywords_when_embedding_fails(session, monkeypatch):
    import retrieval.hybrid as hybrid

    def boom(*a, **k):
        raise RuntimeError("embedding service unavailable")

    monkeypatch.setattr(hybrid, "embed", boom)
    hits = search(session, "multi-factor authentication for administrative access", k=3)
    assert hits and hits[0].section.citation == "SEC-ACCESS §2.2"
    assert all(h.vector_rank is None for h in hits)
